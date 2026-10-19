"""Tests for the official-form fill / validation pipeline.

The pipeline cannot be exercised against the real templates (they are not on this
machine), so these tests build a **synthetic fixture** with python-docx: a table of
six labelled section rows plus a few labelled metadata rows. The fixture is named
and built as a fixture, never as an official template -- and the fill script's
"refuse to treat our own draft as official" rule is tested by putting this
repository's own ``非官方模板`` banner into one of them.

Both scripts import python-docx lazily, so the whole module skips when it is absent.
"""

from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
GRANT_SCRIPTS = REPO / "scripts" / "grant"
SIX = REPO / "docs" / "teaching_grant_2026" / "evidence" / "E5_protocol" / "E5_04_application_text.md"

docx = pytest.importorskip("docx", reason="python-docx is needed for the DOCX pipeline tests")
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

if str(GRANT_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(GRANT_SCRIPTS))


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, GRANT_SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


fill = _load("fill_official_forms")
validate = _load("validate_official_docx")


# --------------------------------------------------------------------------- fixtures


def _cant_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _chinese_normal_style(document) -> None:
    style = document.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(10.5)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "宋体")


def _margins(document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)


def _clear_margins(document) -> None:
    """Simulate an official template that inherits its margins (all unset)."""
    section = document.sections[0]
    for attribute in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attribute, None)


def fixture_docx(path: Path, *, banner: bool = False, drop: str | None = None,
                 prefilled: bool = False, rows: int = 6, margins: bool = True,
                 chinese_style: bool = True, cant_split: bool = True) -> Path:
    """A SYNTHETIC stand-in for 附件2 -- field labels only, no official styling.

    This is a test fixture. It is not, and must never be presented as, the official
    template: the real one is supplied by the applicant and is not on this machine.
    """
    document = docx.Document()
    if chinese_style:
        _chinese_normal_style(document)
    if margins:
        _margins(document)
    if banner:
        document.add_paragraph("非官方模板 · 可编辑草稿 · 测试夹具")
    table = document.add_table(rows=rows, cols=2)
    table.style = "Table Grid"
    placed = 0
    for name, _limit in fill.apptext.SECTIONS:
        if name == drop or placed >= rows:
            continue
        table.cell(placed, 0).text = name
        if prefilled:
            table.cell(placed, 1).text = "已有内容"
        if cant_split:
            _cant_split(table.rows[placed])
        placed += 1
    meta = document.add_table(rows=2, cols=2)
    meta.style = "Table Grid"
    for index, (label, value) in enumerate((("项目名称", ""), ("项目类别", ""))):
        meta.cell(index, 0).text = label
        meta.cell(index, 1).text = value
        if cant_split:
            _cant_split(meta.rows[index])
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def fixture_summary_docx(path: Path, *, banner: bool = False) -> Path:
    document = docx.Document()
    _chinese_normal_style(document)
    _margins(document)
    if banner:
        document.add_paragraph("非官方模板 · 可编辑草稿 · 测试夹具")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for index, (label, value) in enumerate((("项目名称", ""), ("项目主持人", ""), ("所在单位", ""))):
        table.cell(index, 0).text = label
        table.cell(index, 1).text = value
        _cant_split(table.rows[index])
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


@pytest.fixture()
def map_path(tmp_path: Path) -> Path:
    """Minimal field map: six sections plus two fields, labels only."""
    payload = {
        "attachment2": {
            "sections": [{"key": name, "limit": limit, "labels": [name]}
                         for name, limit in fill.apptext.SECTIONS],
            "fields": [{"key": "项目名称", "value": "测试标题", "labels": ["项目名称"]}],
        },
        "attachment3": {
            "sections": [],
            "fields": [{"key": "项目名称", "value": "测试标题", "labels": ["项目名称"]}],
        },
    }
    path = tmp_path / "map.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


# ------------------------------------------------------------------------ source text


def test_six_section_text_parses_from_the_real_file():
    texts = fill.section_texts(SIX)
    assert [name for name, _ in fill.apptext.SECTIONS] == list(texts)
    for name, lines in texts.items():
        assert lines, name
        # Declaration lines and rules are structural, never part of the body.
        assert not any(line.startswith("> 字数") for line in lines), name


def test_section_counts_agree_with_the_application_text_checker():
    checker = _load("check_application_text")
    report = checker.check(SIX)
    texts = fill.section_texts(SIX)
    for row in report["sections"]:
        assert sum(1 for ch in "".join(texts[row["section"]]) if not ch.isspace()) == row["count"]


# ---------------------------------------------------------------------------- refusing


def test_refuses_when_official_files_are_absent(tmp_path, capsys):
    out = tmp_path / "out"
    code = fill.main([
        "--attachment2", str(tmp_path / "missing2.docx"),
        "--attachment3", str(tmp_path / "missing3.docx"),
        "--out-dir", str(out),
    ])
    assert code == 2
    assert not out.exists(), "nothing may be generated when the official files are absent"
    assert "拒绝运行" in capsys.readouterr().err


def test_absence_message_lists_every_accepted_location(tmp_path, capsys):
    # The message must tell the applicant every way to supply the files, including
    # the ZIP, so following the document cannot dead-end at exit 2.
    code = fill.main([
        "--zip", str(tmp_path / "nope.zip"),
        "--official-dir", str(tmp_path / "empty"),
        "--out-dir", str(tmp_path / "out"),
    ])
    assert code == 2
    err = capsys.readouterr().err
    assert "--zip" in err and "--official-dir" in err
    assert "--attachment2" in err and "--attachment3" in err
    assert str(tmp_path / "nope.zip") in err
    assert str(tmp_path / "empty") in err


# ------------------------------------------------------------------------------- ZIP


def _official_zip(tmp_path: Path, *, member2: str, member3: str,
                  decoys: bool = True) -> Path:
    """A ZIP shaped like the applicant's: a subdirectory, plus decoys to ignore."""
    two = fixture_docx(tmp_path / "src2.docx")
    three = fixture_summary_docx(tmp_path / "src3.docx")
    archive = tmp_path / "official.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.write(two, member2)
        bundle.write(three, member3)
        if decoys:
            bundle.writestr("__MACOSX/._附件2.docx", b"junk")
            bundle.writestr("~$附件2.docx", b"lock file")
            bundle.writestr("材料/附件1：项目选题指南.pdf", b"%PDF-1.4 decoy")
            bundle.writestr("材料/申报通知.pdf", b"%PDF-1.4 decoy")
    return archive


@pytest.fixture()
def minimal_map(tmp_path: Path) -> Path:
    path = tmp_path / "map.json"
    path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": name, "limit": limit, "labels": [name]}
                                     for name, limit in fill.apptext.SECTIONS], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    return path


def test_fills_directly_from_the_zip_the_document_tells_the_applicant_to_drop(tmp_path, minimal_map):
    archive = _official_zip(
        tmp_path,
        member2="2026年浙江省教改项目申报材料/附件2：教学改革项目申报书.docx",
        member3="2026年浙江省教改项目申报材料/附件3：申报汇总表.docx",
    )
    out = tmp_path / "out"
    code = fill.main(["--zip", str(archive), "--out-dir", str(out),
                      "--field-map", str(minimal_map)])
    assert code == 0
    produced = docx.Document(str(out / fill.OUT_ATTACHMENT2.name))
    for name, _limit in fill.apptext.SECTIONS:
        assert fill.locate(produced, [name]) is not None, name
    assert (out / fill.OUT_ATTACHMENT3.name).is_file()


def test_zip_members_are_reported_and_decoys_ignored(tmp_path, capsys, minimal_map):
    archive = _official_zip(
        tmp_path,
        member2="材料/附件2：教学改革项目申报书.docx",
        member3="材料/附件3：申报汇总表.docx",
    )
    fill.main(["--zip", str(archive), "--out-dir", str(tmp_path / "out"),
               "--field-map", str(minimal_map)])
    printed = capsys.readouterr().out
    assert "附件2：教学改革项目申报书.docx" in printed
    assert "附件3：申报汇总表.docx" in printed
    assert "未解压" in printed, "the notice and 附件1 must be reported as not extracted"


def test_zip_contents_never_land_in_the_repository(tmp_path, minimal_map):
    archive = _official_zip(
        tmp_path,
        member2="材料/附件2：教学改革项目申报书.docx",
        member3="材料/附件3：申报汇总表.docx",
    )
    out = tmp_path / "out"
    before = {p.name for p in fill.OFFICIAL_DIR.iterdir()}
    assert fill.main(["--zip", str(archive), "--out-dir", str(out),
                      "--field-map", str(minimal_map)]) == 0
    after = {p.name for p in fill.OFFICIAL_DIR.iterdir()}
    assert before == after, "extraction must not write into the repo's official/ directory"
    # And the temporary extraction directory is gone once the run finishes.
    leftovers = list(Path(tempfile.gettempdir()).glob("clinpath-official-*"))
    assert leftovers == [], leftovers


def test_zip_with_only_one_form_still_reports_the_other_as_missing(tmp_path, capsys, minimal_map):
    two = fixture_docx(tmp_path / "src2.docx")
    archive = tmp_path / "partial.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.write(two, "材料/附件2：教学改革项目申报书.docx")
    code = fill.main(["--zip", str(archive), "--out-dir", str(tmp_path / "out"),
                      "--field-map", str(minimal_map)])
    assert code == 2
    assert "attachment3" in capsys.readouterr().err


def test_unreadable_suffix_inside_the_zip_is_explained(tmp_path, capsys, minimal_map):
    archive = tmp_path / "doc-only.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr("材料/附件2：教学改革项目申报书.doc", b"\xd0\xcf\x11\xe0 old format")
        bundle.writestr("材料/附件3：申报汇总表.doc", b"\xd0\xcf\x11\xe0 old format")
    code = fill.main(["--zip", str(archive), "--out-dir", str(tmp_path / "out"),
                      "--field-map", str(minimal_map)])
    assert code == 2
    err = capsys.readouterr().err
    # The applicant put the ZIP where the document said to; the failure must say
    # what is actually wrong with it, not just that nothing was found.
    assert "不是 .docx" in err
    assert "已找到的内容" in err


def test_zip_member_matching_prefers_the_right_attachment():
    hints2 = fill.MEMBER_HINTS["attachment2"]
    hints3 = fill.MEMBER_HINTS["attachment3"]
    names = ["材料/附件2：教学改革项目申报书.docx", "材料/附件3：申报汇总表.docx"]
    assert Path(fill.best_match(names, hints2)).name.startswith("附件2")
    assert Path(fill.best_match(names, hints3)).name.startswith("附件3")
    # A lock file or a resource-fork entry must never win.
    noisy = ["~$附件2.docx", "__MACOSX/._附件2.docx", "材料/附件2：教学改革项目申报书.docx"]
    assert Path(fill.best_match(noisy, hints2)).name == "附件2：教学改革项目申报书.docx"


def test_extracted_docx_in_the_official_dir_is_also_accepted(tmp_path, minimal_map):
    official = tmp_path / "official"
    official.mkdir()
    fixture_docx(official / "附件2：教学改革项目申报书.docx")
    fixture_summary_docx(official / "附件3：申报汇总表.docx")
    out = tmp_path / "out"
    code = fill.main(["--official-dir", str(official), "--out-dir", str(out),
                      "--field-map", str(minimal_map)])
    assert code == 0
    assert (out / fill.OUT_ATTACHMENT2.name).is_file()


def test_explicit_paths_win_over_zip_and_directory(tmp_path, capsys, minimal_map):
    archive = _official_zip(tmp_path, member2="材料/附件2：教学改革项目申报书.docx",
                            member3="材料/附件3：申报汇总表.docx")
    explicit2 = fixture_docx(tmp_path / "explicit2.docx")
    explicit3 = fixture_summary_docx(tmp_path / "explicit3.docx")
    out = tmp_path / "out"
    assert fill.main(["--zip", str(archive), "--attachment2", str(explicit2),
                      "--attachment3", str(explicit3), "--out-dir", str(out),
                      "--field-map", str(minimal_map)]) == 0
    printed = capsys.readouterr().out
    assert "命令行指定" in printed


def test_the_home_config_directory_is_really_searched(tmp_path, monkeypatch, minimal_map):
    """Path B in the document must be a location the script actually looks in."""
    home_official = tmp_path / "config" / "clinpath" / "official"
    home_official.mkdir(parents=True)
    archive = _official_zip(tmp_path, member2="材料/附件2：教学改革项目申报书.docx",
                            member3="材料/附件3：申报汇总表.docx")
    # Move the ZIP to the "home" location and point the search list at it.
    target = home_official / fill.ZIP_NAME
    target.write_bytes(archive.read_bytes())
    monkeypatch.setattr(fill, "OFFICIAL_DIRS", (tmp_path / "absent", home_official))
    monkeypatch.setattr(fill, "ZIP_CANDIDATES", (tmp_path / "absent" / fill.ZIP_NAME, target))

    out = tmp_path / "out"
    assert fill.main(["--official-dir", str(tmp_path / "absent"), "--out-dir", str(out),
                      "--field-map", str(minimal_map)]) == 0
    assert (out / fill.OUT_ATTACHMENT2.name).is_file()


def test_a_bad_explicit_zip_still_falls_back_to_the_default_locations(
        tmp_path, monkeypatch, minimal_map):
    home_official = tmp_path / "config" / "clinpath" / "official"
    home_official.mkdir(parents=True)
    archive = _official_zip(tmp_path, member2="材料/附件2：教学改革项目申报书.docx",
                            member3="材料/附件3：申报汇总表.docx")
    target = home_official / fill.ZIP_NAME
    target.write_bytes(archive.read_bytes())
    monkeypatch.setattr(fill, "OFFICIAL_DIRS", (tmp_path / "absent", home_official))
    monkeypatch.setattr(fill, "ZIP_CANDIDATES", (target,))

    out = tmp_path / "out"
    assert fill.main(["--zip", str(tmp_path / "typo.zip"), "--official-dir",
                      str(tmp_path / "absent"), "--out-dir", str(out),
                      "--field-map", str(minimal_map)]) == 0
    assert (out / fill.OUT_ATTACHMENT2.name).is_file()


def test_refuses_to_fill_our_own_generated_draft(tmp_path, capsys):
    two = fixture_docx(tmp_path / "a2.docx", banner=True)
    three = fixture_summary_docx(tmp_path / "a3.docx", banner=True)
    out = tmp_path / "out"
    code = fill.main(["--attachment2", str(two), "--attachment3", str(three), "--out-dir", str(out)])
    assert code == 6
    assert not out.exists()
    assert "非官方模板" in capsys.readouterr().out


def test_missing_field_is_reported_not_skipped(tmp_path, capsys):
    two = fixture_docx(tmp_path / "a2.docx", drop="项目特色与创新")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {
            "sections": [{"key": name, "limit": limit, "labels": [name]}
                         for name, limit in fill.apptext.SECTIONS],
            "fields": [],
        },
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "out"
    code = fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path)])
    assert code == 3
    printed = capsys.readouterr().out
    assert "未找到 section「项目特色与创新」" in printed
    assert not (out / fill.OUT_ATTACHMENT2.name).exists(), "no partial file without --write-partial"


def test_write_partial_still_writes_what_matched(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx", drop="项目特色与创新")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {
            "sections": [{"key": name, "limit": limit, "labels": [name]}
                         for name, limit in fill.apptext.SECTIONS],
            "fields": [],
        },
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "out"
    code = fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path), "--write-partial"])
    assert code == 3
    produced = out / fill.OUT_ATTACHMENT2.name
    assert produced.is_file()
    assert docx.Document(str(produced)) is not None


def test_existing_content_is_a_conflict_unless_overwrite(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx", prefilled=True)
    three = fixture_summary_docx(tmp_path / "a3.docx")
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": "项目研究基础", "limit": 600,
                                      "labels": ["项目研究基础"]}], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "out"
    args = ["--attachment2", str(two), "--attachment3", str(three),
            "--out-dir", str(out), "--field-map", str(map_path)]
    assert fill.main(args) == 7
    assert not (out / fill.OUT_ATTACHMENT2.name).exists()
    assert fill.main(args + ["--overwrite"]) == 0
    assert (out / fill.OUT_ATTACHMENT2.name).is_file()


# ------------------------------------------------------------------------------ filling


def test_fills_six_sections_and_fields(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    out = tmp_path / "out"
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {
            "sections": [{"key": name, "limit": limit, "labels": [name]}
                         for name, limit in fill.apptext.SECTIONS],
            "fields": [{"key": "项目名称", "value": "测试标题", "labels": ["项目名称"]}],
        },
        "attachment3": {"sections": [], "fields": [
            {"key": "项目名称", "value": "测试标题", "labels": ["项目名称"]}]},
    }, ensure_ascii=False), encoding="utf-8")

    assert fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path)]) == 0

    produced = docx.Document(str(out / fill.OUT_ATTACHMENT2.name))
    texts = fill.section_texts(SIX)
    for name, _limit in fill.apptext.SECTIONS:
        hit = fill.locate(produced, [name])
        assert hit is not None, name
        assert hit["target"].text.strip() == "\n".join(texts[name]) or \
            hit["target"].text.strip() == "".join(texts[name]), name
    assert fill.locate(produced, ["项目名称"])["target"].text.strip() == "测试标题"

    summary = docx.Document(str(out / fill.OUT_ATTACHMENT3.name))
    assert fill.locate(summary, ["项目名称"])["target"].text.strip() == "测试标题"


def test_output_character_counts_match_the_source(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    out = tmp_path / "out"
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": name, "limit": limit, "labels": [name]}
                                     for name, limit in fill.apptext.SECTIONS], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    assert fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path)]) == 0

    produced = docx.Document(str(out / fill.OUT_ATTACHMENT2.name))
    texts = fill.section_texts(SIX)
    for name, _limit in fill.apptext.SECTIONS:
        target = fill.locate(produced, [name])["target"].text
        filled = sum(1 for ch in target if not ch.isspace())
        source = sum(1 for ch in "".join(texts[name]) if not ch.isspace())
        assert filled == source, (name, filled, source)


def test_dry_run_writes_nothing(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    out = tmp_path / "out"
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": name, "limit": limit, "labels": [name]}
                                     for name, limit in fill.apptext.SECTIONS], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    assert fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path), "--dry-run"]) == 0
    assert not out.exists()


def test_over_limit_section_is_refused(tmp_path):
    # A 5-row fixture cannot take six sections, but an over-long *value* is what
    # the limit rule guards: shrink a section limit to 10 characters.
    two = fixture_docx(tmp_path / "a2.docx")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": "项目研究基础", "limit": 10,
                                      "labels": ["项目研究基础"]}], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "out"
    assert fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path)]) == 4
    assert not (out / fill.OUT_ATTACHMENT2.name).exists()


def test_labels_are_matched_with_numbering_and_spacing(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx")
    document = docx.Document(str(two))
    table = document.tables[0]
    table.cell(0, 0).text = "一、项目研究基础"
    table.cell(1, 0).text = "拟解决的重点问题："
    document.save(str(two))

    reopened = docx.Document(str(two))
    assert fill.locate(reopened, ["项目研究基础"]) is not None
    assert fill.locate(reopened, ["拟解决的重点问题"]) is not None


def test_locate_falls_back_to_paragraphs(tmp_path):
    path = tmp_path / "p.docx"
    document = docx.Document()
    document.add_paragraph("项目名称：________________")
    document.add_paragraph("")
    document.add_paragraph("项目周期")
    document.add_paragraph("")
    document.save(str(path))

    reopened = docx.Document(str(path))
    name = fill.locate(reopened, ["项目名称"])
    period = fill.locate(reopened, ["项目周期"])
    assert name["kind"] == "paragraph-append"
    assert period["kind"] == "paragraph-next"

    fill.apply_write(name, ["测试标题"])
    fill.apply_write(period, ["【待填：周期】"])
    assert "测试标题" in reopened.paragraphs[0].text
    assert "____" not in reopened.paragraphs[0].text
    assert reopened.paragraphs[3].text == "【待填：周期】"


# --------------------------------------------------------------------------- validation


def _filled(tmp_path: Path) -> Path:
    two = fixture_docx(tmp_path / "a2.docx")
    three = fixture_summary_docx(tmp_path / "a3.docx")
    out = tmp_path / "out"
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": name, "limit": limit, "labels": [name]}
                                     for name, limit in fill.apptext.SECTIONS], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    assert fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path)]) == 0
    return out / fill.OUT_ATTACHMENT2.name


def _status(report: dict, name: str) -> str:
    return next(item["status"] for item in report["checks"] if item["name"] == name)


def test_validation_passes_a_well_formed_fixture(tmp_path):
    produced = _filled(tmp_path)
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "文件与重新打开") == "PASS"
    assert _status(report, "字数核验") == "PASS"
    assert _status(report, "页边距") == "PASS"
    assert _status(report, "表格完整性") == "PASS"
    assert _status(report, "跨页：表格内不得有显式分页符") == "PASS"
    assert _status(report, "跨页：多行表格设置行内不拆分") == "PASS"
    assert _status(report, "中文字体与字号") == "PASS"
    # Without the official template, agreement cannot be claimed.
    assert _status(report, "页边距与官方模板一致") == "SKIPPED"


def test_validation_reports_a_missing_file(tmp_path):
    assert validate.main([str(tmp_path / "nope.docx")]) == 2


def test_validation_flags_unset_chinese_fonts(tmp_path):
    produced = _filled(tmp_path)
    document = docx.Document(str(produced))
    document.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts().attrib.clear()
    document.styles["Normal"].font.size = None
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "中文字体与字号") == "FAIL"


def test_validation_flags_a_ragged_table(tmp_path):
    produced = _filled(tmp_path)
    document = docx.Document(str(produced))
    table = document.tables[0]
    table.rows[0]._tr.remove(table.rows[0].cells[-1]._tc)
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "表格完整性") == "FAIL"


def test_validation_flags_an_explicit_page_break_inside_a_table(tmp_path):
    produced = _filled(tmp_path)
    document = docx.Document(str(produced))
    run = document.tables[0].cell(0, 0).paragraphs[0].add_run()
    run.add_break(docx.enum.text.WD_BREAK.PAGE)
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "跨页：表格内不得有显式分页符") == "FAIL"


def test_validation_flags_rows_that_may_split(tmp_path):
    produced = _filled(tmp_path)
    document = docx.Document(str(produced))
    for row in document.tables[0].rows:
        properties = row._tr.trPr
        if properties is None:
            continue
        for element in properties.findall(qn("w:cantSplit")):
            properties.remove(element)
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "跨页：多行表格设置行内不拆分") == "FAIL"


def test_validation_flags_missing_margins(tmp_path):
    two = fixture_docx(tmp_path / "a2.docx", margins=False)
    three = fixture_summary_docx(tmp_path / "a3.docx")
    out = tmp_path / "out"
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps({
        "attachment2": {"sections": [{"key": name, "limit": limit, "labels": [name]}
                                     for name, limit in fill.apptext.SECTIONS], "fields": []},
        "attachment3": {"sections": [], "fields": []},
    }, ensure_ascii=False), encoding="utf-8")
    assert fill.main(["--attachment2", str(two), "--attachment3", str(three),
                      "--out-dir", str(out), "--field-map", str(map_path)]) == 0
    produced = out / fill.OUT_ATTACHMENT2.name
    document = docx.Document(str(produced))
    _clear_margins(document)
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "页边距") == "FAIL"


def test_validation_compares_margins_with_the_template(tmp_path):
    produced = _filled(tmp_path)
    template = fixture_docx(tmp_path / "a2.docx")
    report = validate.validate(produced, template, fill.section_texts(SIX), docx)
    assert _status(report, "页边距与官方模板一致") == "PASS"

    shifted = docx.Document(str(template))
    shifted.sections[0].left_margin = Cm(1.0)
    shifted_path = tmp_path / "shifted.docx"
    shifted.save(str(shifted_path))
    report = validate.validate(produced, shifted_path, fill.section_texts(SIX), docx)
    assert _status(report, "页边距与官方模板一致") == "FAIL"


def test_validation_reports_over_limit_and_placeholder(tmp_path):
    produced = _filled(tmp_path)
    document = docx.Document(str(produced))
    target = fill.locate(document, ["项目特色与创新"])["target"]
    target.text = "字" * 400
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "字数核验") == "FAIL"

    document = docx.Document(str(produced))
    fill.locate(document, ["项目名称"])["target"].text = "【待填：姓名】"
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "提交就绪：无【待填】占位符") == "FAIL"


def test_placeholder_check_passes_only_once_the_text_has_no_markers(tmp_path):
    # The real E5_04 body deliberately carries 13 【待填】 markers, so a fill from
    # the real source must FAIL submission-readiness -- that is the honest result,
    # not a bug. Removing the markers is what a human filling them in looks like.
    produced = _filled(tmp_path)
    assert _status(validate.validate(produced, None, fill.section_texts(SIX), docx),
                   "提交就绪：无【待填】占位符") == "FAIL"

    document = docx.Document(str(produced))
    for paragraph in fill.iter_paragraphs(document):
        for run in paragraph.runs:
            run.text = run.text.replace("【待填", "已填")
    document.save(str(produced))
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    assert _status(report, "提交就绪：无【待填】占位符") == "PASS"


def test_pdf_render_check_is_skipped_honestly_when_no_renderer(tmp_path, monkeypatch):
    monkeypatch.setattr(validate.shutil, "which", lambda _name: None)
    produced = _filled(tmp_path)
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    item = next(item for item in report["checks"] if item["name"] == "跨页：PDF 渲染目视检查")
    assert item["status"] == "SKIPPED"
    # Fixed wording: the report must never drift into implying the check ran.
    assert item["detail"] == (
        "no libreoffice/pandoc on this machine; requires a human render check")
    assert "libreoffice --headless --convert-to pdf" in item["human_command"]
    # A SKIPPED check is never counted as a pass -- the report must say so.
    text = validate.render(report)
    assert "未能执行的检查" in text
    assert "no libreoffice/pandoc on this machine" in text
    assert not any(item["status"] == "PASS" for item in report["checks"]
                   if item["name"] == "跨页：PDF 渲染目视检查")


def test_no_renderer_reason_is_the_agreed_wording():
    assert validate.NO_RENDERER_REASON == (
        "no libreoffice/pandoc on this machine; requires a human render check")
    command = validate.render_command(Path("out.docx"))
    assert command.startswith("libreoffice --headless --convert-to pdf")
    assert "out.docx" in command


def test_render_check_does_not_run_when_no_renderer_is_on_path(tmp_path, monkeypatch):
    """Even if a renderer name were found, absence of the binary must not 'pass'."""
    monkeypatch.setattr(validate.shutil, "which", lambda _name: None)

    def explode(*_args, **_kwargs):  # pragma: no cover - must never be called
        raise AssertionError("the renderer must not be invoked when it is absent")

    monkeypatch.setattr(validate.subprocess, "run", explode, raising=False)
    produced = _filled(tmp_path)
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    item = next(item for item in report["checks"] if item["name"] == "跨页：PDF 渲染目视检查")
    assert item["status"] == "SKIPPED"


def test_render_states_skipped_checks_are_not_passes(tmp_path, monkeypatch):
    monkeypatch.setattr(validate.shutil, "which", lambda _name: None)
    produced = _filled(tmp_path)
    report = validate.validate(produced, None, fill.section_texts(SIX), docx)
    text = validate.render(report)
    assert "未能执行的检查" in text
    assert "PDF 渲染" in text
