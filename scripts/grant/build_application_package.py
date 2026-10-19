#!/usr/bin/env python3
"""Assemble the R7.3 application package and render Word (.docx) versions.

Inputs (all Markdown, edited by hand):
  docs/teaching_grant_2026/evidence/E5_protocol/E5_04_application_text.md  six-section source of truth
  docs/teaching_grant_2026/APPLICATION_PACKAGE/{10,20,30,40,50,60,70,80,90}_*.md

Outputs:
  APPLICATION_PACKAGE/01_FULL_APPLICATION_DRAFT.md   generated: basic info + six sections + appendices
  APPLICATION_PACKAGE/docx/*.docx                    one Word file per source Markdown file

The Word files are **not** the official template: the applicant's official originals
(`ClinPath_2026_Official_Source_Templates.zip`, restored in the ChatGPT file sandbox)
are not readable from this machine -- see
`docs/teaching_grant_2026/FINAL_DELIVERY/OFFICIAL_APPLICATION_VALIDATION.md` -- so
every file carries that banner. python-docx is needed only for the .docx step:

    uv run --python 3.11 --with python-docx python scripts/grant/build_application_package.py
    python3 scripts/grant/build_application_package.py --markdown-only     # no python-docx needed

After running, re-check the text with::

    python3 scripts/grant/check_application_text.py
    python3 scripts/grant/check_protocol_consistency.py
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
GRANT = ROOT / "docs" / "teaching_grant_2026"
PACKAGE = GRANT / "APPLICATION_PACKAGE"
SIX = GRANT / "evidence" / "E5_protocol" / "E5_04_application_text.md"
TITLE = "基于人工智能辅助临床推理训练的医学生个性化学习与形成性评价教学改革研究"
BANNER = "非官方模板 · 可编辑草稿 · 含大量【待填】 · 未经任何真实人员审核或签章"

APPENDICES = [
    "10_basic_info_and_host_CV.md",
    "20_team_roles.md",
    "30_implementation_plan.md",
    "40_outcomes_and_KPIs.md",
    "50_budget_suggestion.md",
    "60_unit_support_checklist.md",
    "70_summary_table_attachment3_draft.md",
    "80_one_page_summary.md",
    "90_expert_demo_note.md",
]


def demote(text: str, levels: int = 1) -> str:
    """Push every Markdown heading down so appendices never collide with the six sections."""
    out = []
    for line in text.splitlines():
        match = re.match(r"^(#{1,5})(\s+.*)$", line)
        out.append("#" * (len(match.group(1)) + levels) + match.group(2) if match else line)
    return "\n".join(out)


def six_sections() -> str:
    text = SIX.read_text(encoding="utf-8")
    start = text.index("## 一、项目研究基础")
    end = text.index("## 字数校验")
    return text[start:end].rstrip().rstrip("-").rstrip()


def build_full_draft() -> str:
    parts = [
        f"# 《{TITLE}》申报书（完整可编辑草稿）\n",
        f"> {BANNER}\n>\n"
        "> **GENERATED**：本文件由 `scripts/grant/build_application_package.py` 拼装，请改源文件而不是本文件。\n"
        "> 官方模板原件未获得（见 `00_README.md`）；六栏名称与字数上限取自任务书，转入官方 Word 后须再核对。\n",
        "---\n",
        "## 〇、项目基本信息与主持人教学简历\n",
        demote(
            (PACKAGE / APPENDICES[0]).read_text(encoding="utf-8").split("\n", 1)[1], 1
        ),
        "\n---\n",
        six_sections(),
        "\n---\n",
        "## 附录\n",
    ]
    for name in APPENDICES[1:7]:
        parts.append(demote((PACKAGE / name).read_text(encoding="utf-8"), 2))
        parts.append("\n")
    return "\n".join(parts).rstrip() + "\n"


# --------------------------------------------------------------------------- docx


def _add_runs(paragraph, text: str) -> None:
    """Render **bold** and `code` spans; everything else as plain text."""
    for chunk in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            paragraph.add_run(chunk[1:-1])
        else:
            paragraph.add_run(chunk)


def markdown_to_docx(markdown: str, out: Path) -> None:
    from docx import Document  # imported lazily: only needed for the .docx step
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(10.5)

    banner = doc.add_paragraph()
    run = banner.add_run(BANNER)
    run.bold = True

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            doc.add_heading(re.sub(r"\*\*", "", heading.group(2)), level=min(len(heading.group(1)), 4))
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            _add_runs(paragraph, stripped.lstrip("> ").strip())
            i += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.、]\s+(.*)$", stripped)
        if bullet or numbered:
            paragraph = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            _add_runs(paragraph, (bullet or numbered).group(1))
            i += 1
            continue
        paragraph = doc.add_paragraph()
        _add_runs(paragraph, stripped)
        i += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--markdown-only", action="store_true", help="只生成完整草稿 Markdown，不生成 .docx")
    args = parser.parse_args(argv)

    for name in APPENDICES:
        if not (PACKAGE / name).is_file():
            print(f"错误：缺少 {PACKAGE / name}", file=sys.stderr)
            return 2
    if not SIX.is_file():
        print(f"错误：缺少六栏正文 {SIX}", file=sys.stderr)
        return 2

    full = build_full_draft()
    (PACKAGE / "01_FULL_APPLICATION_DRAFT.md").write_text(full, encoding="utf-8")
    print("写入 01_FULL_APPLICATION_DRAFT.md")
    if args.markdown_only:
        return 0

    try:
        import docx  # noqa: F401
    except ModuleNotFoundError:
        print("错误：未安装 python-docx；请用 uv run --with python-docx 运行，或加 --markdown-only", file=sys.stderr)
        return 3

    targets = {
        "01_FULL_APPLICATION_DRAFT": full,
        "02_six_sections": six_sections(),
        **{Path(n).stem: (PACKAGE / n).read_text(encoding="utf-8") for n in APPENDICES},
    }
    for stem, text in targets.items():
        markdown_to_docx(text, PACKAGE / "docx" / f"{stem}.docx")
        print(f"写入 docx/{stem}.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
